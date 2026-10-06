"""
J.A.R.V.I.S. Multi-Format Document Ingestion Engine
Extracts and structures readable text from PDFs, Word documents (.docx),
PowerPoint slides (.pptx), code files, and plain text documents.
"""

import os
from pypdf import PdfReader
import docx
from pptx import Presentation

MAX_EXTRACT_CHARS = 35000  # ~7,000 to 8,000 words (plenty for lecture notes, research papers, assignments)

def extract_document_text(file_path: str, max_chars: int = MAX_EXTRACT_CHARS) -> dict:
    """
    Extracts text from a given file path based on its extension.
    Supports .pdf, .docx, .pptx, .txt, .py, .c, .cpp, .csv, .md, .json.
    """
    if not os.path.exists(file_path):
        return {
            "success": False,
            "error": f"File not found: {file_path}",
            "text": ""
        }

    filename = os.path.basename(file_path)
    _, ext = os.path.splitext(filename.lower())

    extracted_text = ""
    file_type_label = ext.replace(".", "").upper()

    try:
        # 1. PDF Document (.pdf)
        if ext == '.pdf':
            reader = PdfReader(file_path)
            pages = []
            total_pages = len(reader.pages)
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages.append(f"--- [Page {i+1} of {total_pages}] ---\n{page_text.strip()}")
            extracted_text = "\n\n".join(pages)
            file_type_label = f"PDF Document ({total_pages} pages)"

        # 2. Microsoft Word Document (.docx)
        elif ext == '.docx':
            doc = docx.Document(file_path)
            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())
            
            # Extract tables if present
            for table_idx, table in enumerate(doc.tables):
                table_lines = [f"[Table {table_idx+1}]"]
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_lines.append(" | ".join(row_data))
                if len(table_lines) > 1:
                    paragraphs.append("\n".join(table_lines))

            extracted_text = "\n\n".join(paragraphs)
            file_type_label = "Word Document (.docx)"

        # 3. Microsoft PowerPoint Presentation (.pptx)
        elif ext in ['.pptx', '.ppt']:
            if ext == '.pptx':
                prs = Presentation(file_path)
                slides = []
                total_slides = len(prs.slides)
                for i, slide in enumerate(prs.slides):
                    slide_elements = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                t = paragraph.text.strip()
                                if t:
                                    slide_elements.append(t)
                    if slide_elements:
                        slides.append(f"--- [Slide {i+1} of {total_slides}] ---\n" + "\n".join(slide_elements))
                extracted_text = "\n\n".join(slides)
                file_type_label = f"PowerPoint Presentation ({total_slides} slides)"
            else:
                extracted_text = "Legacy .ppt binary format detected. Please save or export as .pptx for full deep extraction, Sir."

        # 4. Plain Text, Code, & Markdown Files
        elif ext in ['.txt', '.py', '.c', '.cpp', '.h', '.csv', '.md', '.json', '.html', '.css', '.js', '.log', '.env']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                extracted_text = f.read()
            file_type_label = f"Code/Text File ({ext})"

        # 5. Image Files (.png, .jpg, .jpeg)
        elif ext in ['.png', '.jpg', '.jpeg', '.webp']:
            file_size_kb = round(os.path.getsize(file_path) / 1024, 1)
            extracted_text = f"[Image Attached: {filename} ({file_size_kb} KB)]. Multimodal visual file uploaded for reference."
            file_type_label = "Image File"

        else:
            # Fallback text attempt
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    extracted_text = f.read()
                file_type_label = f"Generic Document ({ext})"
            except Exception:
                return {
                    "success": False,
                    "error": f"Unsupported document format: {ext}",
                    "text": ""
                }

        # Check truncation
        is_truncated = False
        original_length = len(extracted_text)
        if len(extracted_text) > max_chars:
            extracted_text = extracted_text[:max_chars] + f"\n\n... [Content truncated for memory optimization: showing first {max_chars} of {original_length} characters] ..."
            is_truncated = True

        word_count = len(extracted_text.split())

        return {
            "success": True,
            "filename": filename,
            "file_type": file_type_label,
            "word_count": word_count,
            "text": extracted_text,
            "truncated": is_truncated
        }

    except Exception as e:
        return {
            "success": False,
            "error": f"Failed to extract document contents: {str(e)}",
            "text": ""
        }
